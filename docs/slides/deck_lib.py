"""Small helpers on top of python-pptx for the project deck (theme, text, shapes, arrows, tables, charts)."""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------- theme
INK = "0F2A33"      # dominant dark (title / section slides, body text)
TEAL = "1B4D5C"     # secondary dark
PALE = "EAF2F3"     # light tint for cards
WHITE = "FFFFFF"
AMBER = "F2A541"    # the "spike" accent: badges, highlights
GREEN = "1B998B"    # benign
RED = "D1495B"      # malicious
BLUE = "2E86AB"
PURPLE = "7A5C99"
GREY = "5F6F78"     # muted text
LINE = "C9D6DA"
HEAD = "Cambria"
BODY = "Calibri"
W, H = 13.333, 7.5


def rgb(h):
    return RGBColor.from_string(h)


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


class Deck:
    def __init__(self):
        self.prs = new_deck()
        self.n = 0

    def slide(self, dark=False, notes=None, section=None):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(INK if dark else WHITE)
        self.n += 1
        if self.n > 1:
            text(s, W - 1.3, H - 0.5, 0.9, 0.3, str(self.n), 11, GREY if not dark else "9FB3BA", align="r")
            if section and not dark:
                text(s, 0.6, H - 0.5, 8, 0.3, section, 11, GREY)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    def save(self, path):
        self.prs.save(path)


# ---------------------------------------------------------------- text
def _para_fmt(p, size, color, bold=False, italic=False, font=BODY, align="l"):
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(color)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font


def text(s, x, y, w, h, content, size=16, color=INK, bold=False, italic=False, font=BODY, align="l",
         anchor="t", margin=0.05, name=None, space_after=4):
    """content: str (\\n = new paragraph) or list of paragraphs; a paragraph is a str or a list of
    (text, {bold, color, italic, size}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(margin))
    paras = content.split("\n") if isinstance(content, str) else content
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
        _para_fmt(p, size, color, bold, italic, font, align)
        for r, (t, o) in zip(p.runs, runs):
            if "bold" in o: r.font.bold = o["bold"]
            if "color" in o: r.font.color.rgb = rgb(o["color"])
            if "italic" in o: r.font.italic = o["italic"]
            if "size" in o: r.font.size = Pt(o["size"])
            if "font" in o: r.font.name = o["font"]
    return tb


def bullets(s, x, y, w, h, items, size=16, color=INK, bullet_color=AMBER, space_after=8, name=None):
    """items: str or list of runs (see text()); a (level, item) tuple indents one level."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(0.05))
    for i, it in enumerate(items):
        lvl = 0
        if isinstance(it, tuple):
            lvl, it = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        runs = [(it, {})] if isinstance(it, str) else it
        for t, o in runs:
            p.add_run().text = t
        _para_fmt(p, size - 2 * lvl, color)
        for r, (t, o) in zip(p.runs, runs):
            if "bold" in o: r.font.bold = o["bold"]
            if "color" in o: r.font.color.rgb = rgb(o["color"])
            if "italic" in o: r.font.italic = o["italic"]
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.3 + 0.3 * lvl))))
        pPr.set("indent", str(-int(Inches(0.24))))
        buClr = pPr.makeelement(qn("a:buClr"), {})
        clr = buClr.makeelement(qn("a:srgbClr"), {"val": bullet_color})
        buClr.append(clr)
        pPr.append(buClr)
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•" if lvl == 0 else "–"}))
    return tb


def title(s, t, dark=False, sub=None):
    text(s, 0.6, 0.35, 12.1, 0.85, t, 34, WHITE if dark else INK, bold=True, font=HEAD, anchor="m", name="Title")
    if sub:
        text(s, 0.6, 1.15, 12.1, 0.45, sub, 16, "B9CBD0" if dark else GREY, italic=True, name="Subtitle")


# ---------------------------------------------------------------- shapes
def box(s, x, y, w, h, fill=PALE, line=None, shape="round", txt=None, size=15, color=INK, bold=False,
        align="c", anchor="m", radius=0.12, name=None, font=BODY, line_w=1.25):
    kind = {"round": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE, "oval": MSO_SHAPE.OVAL,
            "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON}[shape]
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if shape == "round":
        sh.adjustments[0] = radius
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, Inches(0.08))
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if txt is not None:
        paras = txt.split("\n") if isinstance(txt, str) else txt
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            runs = [(para, {})] if isinstance(para, str) else para
            for t, o in runs:
                p.add_run().text = t
            _para_fmt(p, size, color, bold, False, font, align)
            for r, (t, o) in zip(p.runs, runs):
                if "bold" in o: r.font.bold = o["bold"]
                if "color" in o: r.font.color.rgb = rgb(o["color"])
                if "size" in o: r.font.size = Pt(o["size"])
                if "italic" in o: r.font.italic = o["italic"]
    return sh


def badge(s, x, y, label, fill=AMBER, color=INK, d=0.55, size=18):
    return box(s, x, y, d, d, fill=fill, shape="oval", txt=label, size=size, color=color, bold=True)


def card(s, x, y, w, h, head, body, fill=PALE, head_color=INK, body_size=14, head_size=17, badge_label=None,
         badge_fill=AMBER):
    box(s, x, y, w, h, fill=fill)
    tx = x + 0.25
    if badge_label is not None:
        badge(s, x + 0.2, y + 0.2, badge_label, fill=badge_fill, color=WHITE if badge_fill in (RED, GREEN, BLUE, PURPLE, TEAL, INK) else INK, d=0.5, size=16)
        tx = x + 0.85
    text(s, tx, y + 0.2, w - (tx - x) - 0.2, 0.55, head, head_size, head_color, bold=True, anchor="m")
    text(s, x + 0.25, y + 0.85, w - 0.45, h - 1.0, body, body_size, INK)


def arrow(s, x1, y1, x2, y2, color=TEAL, width=2.25, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": "dash"}))
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return c


def line(s, x1, y1, x2, y2, color=LINE, width=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def picture(s, path, x, y, w=None, h=None, descr=""):
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)
    pic._element.nvPicPr.cNvPr.set("descr", descr)
    return pic


# ---------------------------------------------------------------- tables
def table(s, x, y, w, col_w, rows, size=14, head_fill=TEAL, head_color=WHITE, zebra=PALE, row_h=0.45,
          bold_first_col=True, col_colors=None):
    nr, nc = len(rows), len(rows[0])
    shp = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = shp.table
    tot = sum(col_w)
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(w * cw / tot)
    for r in range(nr):
        t.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = t.cell(r, c)
            val = rows[r][c]
            runs = [(val, {})] if isinstance(val, str) else val
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            for tx, o in runs:
                rr = p.add_run()
                rr.text = tx
                rr.font.size = Pt(size)
                rr.font.name = BODY
                rr.font.bold = (r == 0) or (bold_first_col and c == 0) or o.get("bold", False)
                col = head_color if r == 0 else o.get("color", INK)
                if r > 0 and col_colors and c in col_colors and "color" not in o:
                    col = col_colors[c]
                rr.font.color.rgb = rgb(col)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(head_fill if r == 0 else (zebra if r % 2 == 0 else WHITE))
    return t


# ---------------------------------------------------------------- charts
def _style_axes(ch, size=12):
    for ax in (ch.category_axis, ch.value_axis):
        ax.tick_labels.font.size = Pt(size)
        ax.tick_labels.font.color.rgb = rgb(GREY)
        ax.format.line.color.rgb = rgb(LINE)
    ch.value_axis.major_gridlines.format.line.color.rgb = rgb("E3EAEC")


def line_chart(s, x, y, w, h, cats, series, colors, title_txt=None, size=12):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(size)
    for ser, col in zip(ch.plots[0].series, colors):
        ser.format.line.color.rgb = rgb(col)
        ser.format.line.width = Pt(2.5)
        ser.smooth = False
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb(col)
        ser.marker.format.line.color.rgb = rgb(col)
    _style_axes(ch, size)
    if title_txt:
        ch.has_title = True
        ch.chart_title.text_frame.text = title_txt
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    else:
        ch.has_title = False
    return ch


def bar_chart(s, x, y, w, h, cats, vals, colors, horizontal=True, number_format='0', size=12, title_txt=None):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("value", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    ch = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.vary_by_categories = True
    for i, col in enumerate(colors):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
    plot.has_data_labels = True
    plot.data_labels.number_format = number_format
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(size)
    plot.data_labels.font.color.rgb = rgb(INK)
    _style_axes(ch, size)
    if horizontal:
        ch.category_axis.reverse_order = True
    ch.has_title = bool(title_txt)
    if title_txt:
        ch.chart_title.text_frame.text = title_txt
        ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    return ch


def doughnut(s, x, y, w, h, cats, vals, colors, size=12):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("share", vals)
    ch = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd).chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.RIGHT
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(size)
    plot = ch.plots[0]
    for i, col in enumerate(colors):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
    plot.has_data_labels = True
    plot.data_labels.number_format = '0.0"%"'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.font.size = Pt(size)
    plot.data_labels.font.color.rgb = rgb(WHITE)
    plot.data_labels.font.bold = True
    ch.has_title = False
    return ch
